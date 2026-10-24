"""Small, explicit live smoke test. Uses configured search/AI credits; sends no outreach.

Run: python -X utf8 -m scripts.verify_lead_quality
Reports are local/ignored in exports/quality-audit; credentials are never serialized.
"""
import asyncio
import argparse
import json
import logging
import time
from datetime import datetime, timezone
from pathlib import Path

from dotenv import load_dotenv

load_dotenv()
logging.disable(logging.CRITICAL)

from api.routes.leads import LeadBulkVerifyRequest, verify_leads_for_export
from core.services.production_lead_pipeline import ProductionLeadPipeline
from core.services.lead_quality_gate import apply_quality_gate


def write_report(out):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    reports = []
    for name in ("saas-development", "digital-marketing", "accountants"):
        path = out / f"{name}.json"
        if path.exists():
            reports.append(json.loads(path.read_text(encoding="utf-8")))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Checked leads"
    fields = {
        "Company": "company_name", "Category": "industry", "Location": "location",
        "Contact Person": "contact_person_name", "Published Email": "contact_email",
        "Phone (E.164)": "contact_phone", "Website": "company_website", "Quality Score": "quality_score",
        "Quality Status": "quality_status", "Email Check": "verification_status", "Email Check Detail": "verification_message",
        "Quality Explanation": "acceptance_reason",
    }
    sheet.append(list(fields))
    rows = [lead for report in reports for lead in report.get("leads", [])]
    for lead in rows:
        sheet.append([lead.get(key, "") for key in fields.values()])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="243B73")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = min(60, max(18, max(len(str(cell.value or "")) for cell in col) + 2))
    workbook.save(out / "checked-leads.xlsx")
    lines = ["# Lead quality verification", "", f"Generated: {datetime.now(timezone.utc).isoformat()}", "",
        "Live public-website discovery and email checks; no outreach sent.", "",
        "| Category | Location | Returned | Published emails | Valid phone formats | SMTP-confirmed |",
        "|---|---|---:|---:|---:|---:|"]
    for report in reports:
        leads = report.get("leads", [])
        if not leads:
            continue
        lines.append(f"| {leads[0]['industry']} | {leads[0]['location']} | {len(leads)} | {sum(bool(l.get('contact_email')) for l in leads)} | {sum(l.get('phone_validation_status') == 'valid' for l in leads)} | {sum(bool(l.get('email_smtp_verified')) for l in leads)} |")
    lines += ["", "Phone validation checks numbering format, not whether someone will answer. SMTP timeouts do not establish that an email is invalid or deliverable. These company-level leads are marked Needs_Review; unknown people are left blank.", "",
        "The narrower SaaS-product/Delhi trial returned no qualifying leads. SaaS Development is a separate services category."]
    (out / "REPORT.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


async def replay_original(out):
    """Re-check the user's two original public company URLs without searching."""
    pipeline = ProductionLeadPipeline()
    pipeline._target_location = "Delhi"
    targets = await pipeline._verify_targets([
        {"url": "https://www.chlsoftech.com/saas-developement-company.html", "company_name": "CHL Softech"},
        {"url": "https://www.nethues.com/saas-development", "company_name": "Nethues"},
    ], limit=2)
    targets = await pipeline._scrape_contact_pages_batch(targets)
    leads = [await pipeline._score_one(pipeline._normalized_target_without_extract(target),
             industry="SaaS Development", location="Delhi", use_gemini=False) for target in targets]
    leads, stats = apply_quality_gate([lead for lead in leads if lead])
    result = await verify_leads_for_export(LeadBulkVerifyRequest(leads=leads))
    (out / "original-two-rechecked.json").write_text(json.dumps({**result, "gate": stats}, indent=2), encoding="utf-8")
    print(json.dumps({"original_two": [{key: lead.get(key) for key in (
        "company_name", "contact_email", "contact_phone", "location", "quality_score", "verification_status"
    )} for lead in result["leads"]]}), flush=True)


async def audit(industry, location, out):
    started = time.monotonic()
    print(f"START {industry} / {location}", flush=True)
    pipeline = ProductionLeadPipeline()
    result = await pipeline.run(industry=industry, location=location, target_count=2)
    verified = await verify_leads_for_export(LeadBulkVerifyRequest(leads=result["leads"], max_concurrency=2))
    result["leads"] = verified["leads"]
    result["verification"] = verified["summary"]
    result["elapsed_seconds"] = round(time.monotonic() - started, 1)
    filename = out / (industry.lower().replace(" ", "-") + ".json")
    filename.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    summary = {"category": industry, "location": location, "count": len(result["leads"]),
               "seconds": result["elapsed_seconds"], "report": str(filename),
               "leads": [{key: lead.get(key) for key in (
                   "company_name", "company_website", "contact_email", "contact_phone", "contact_person_name",
                   "location", "quality_score", "quality_status", "verification_status", "rejected",
               )} for lead in result["leads"]]}
    print(json.dumps(summary, ensure_ascii=False), flush=True)
    return summary


async def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--case", action="append", choices=["saas", "development", "marketing", "accountants"])
    parser.add_argument("--original", action="store_true")
    parser.add_argument("--report-only", action="store_true")
    args = parser.parse_args()
    cases = {
        "saas": ("SaaS", "Delhi"), "development": ("SaaS Development", "Delhi"),
        "marketing": ("Digital Marketing", "Mumbai"), "accountants": ("Accountants", "London"),
    }
    out = Path("exports/quality-audit")
    out.mkdir(parents=True, exist_ok=True)
    if args.report_only:
        write_report(out)
        return
    if args.original:
        await replay_original(out)
        return
    results = await asyncio.gather(*(audit(*cases[key], out) for key in (args.case or ["development", "marketing", "accountants"])))
    (out / "summary.json").write_text(json.dumps(results, indent=2), encoding="utf-8")
    write_report(out)


if __name__ == "__main__":
    asyncio.run(main())
