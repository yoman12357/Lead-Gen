"""Build a review workbook only when every discovered company has a manual verdict."""
import argparse
import json
import csv
from pathlib import Path

import phonenumbers
from email_validator import validate_email
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.services.lead_search_policy import domain_key


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True)
    parser.add_argument('--evidence', required=True)
    parser.add_argument('--reviews', required=True)
    parser.add_argument('--history', action='append', help='Diagnostic run directory to include in total audit cost; repeatable')
    args = parser.parse_args()
    root, evidence = Path(args.input), Path(args.evidence)
    manual = json.loads(Path(args.reviews).read_text(encoding='utf-8'))
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = 'Quality summary'
    summary_sheet.append(['Business / Location', 'Requested', 'Raw returned', 'Qualified after review', 'Email', 'Mail server found',
        'Phone', 'Phone format valid', 'Independently supported contact route', 'Direct HTTP rechecks', 'Supplemented official sources', 'Seconds'])
    report_rows, stats = [], []
    all_hosts, all_emails, all_phones = [], [], []
    corrections = []
    market_sheets = {}
    cases = []
    for file in sorted(root.glob('*.json')):
        data = json.loads(file.read_text(encoding='utf-8'))
        if isinstance(data, dict) and isinstance(data.get('result'), dict) and 'leads' in data['result']:
            cases.append((file.stem, data))
    assert cases, 'No saved market runs found'
    requested_total = sum(job['request']['target_count'] for _, job in cases)
    raw_total = sum(len(job['result']['leads']) for _, job in cases)
    headers = ['Company', 'Website', 'Location', 'Email', 'Email publication checked', 'Mail server found', 'Email delivery confirmed',
        'Phone', 'Phone numbering valid', 'Phone publication check', 'Quality verdict', 'Reviewer note',
        'Location source', 'Email source', 'Industry source', 'Location qualification basis', 'Source retrieval', 'Checked at']
    for case, job in cases:
        sheet_name = {'marketing-mumbai': 'Mumbai marketing', 'accountants-london': 'London accountants'}.get(case, case[:31])
        assert job['status'] in {'completed', 'partial'}
        requested = job['request']['target_count']
        leads = job['result']['leads']
        assert len(leads) <= requested
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        raw_sheet = workbook.create_sheet(('Review ' + case)[:31])
        raw_sheet.append(headers)
        case_rows = []
        for index, lead in enumerate(leads, 1):
            host = domain_key(lead['company_website'])
            verdict = manual['companies'][host]  # Missing review must fail, not silently pass.
            company_fit = bool(verdict['category_fit'] and verdict['location_fit'] and verdict['operating_company'])
            observed = json.loads((evidence / f'{case}-{index:02}.json').read_text(encoding='utf-8'))
            assert domain_key(observed['website']) == host
            supplemental = manual.get('supplemental_sources', {}).get(host, {})
            overrides = verdict.get('corrections', {})
            if overrides:
                assert set(overrides) <= {'company_name', 'contact_email', 'contact_phone'}
                assert verdict.get('correction_reason') and supplemental.get('urls')
                lead = dict(lead)
                for field, value in overrides.items():
                    corrections.append([case, host, field, lead.get(field, ''), value, verdict['correction_reason'], ' | '.join(supplemental['urls'])])
                    lead[field] = value
                if 'contact_email' in overrides:
                    lead['email_mx_found'] = False  # Never carry DNS evidence across a changed mailbox.
            email, phone = lead.get('contact_email', ''), lead.get('contact_phone', '')
            email_supported = bool(email) and ((observed['email_reproduced'] is True and observed['email'] == email) or email in supplemental.get('emails', []))
            phone_supported = bool(phone) and ((observed['phone_reproduced'] is True and observed['phone'] == phone) or phone in supplemental.get('phones', []))
            if company_fit:
                assert email_supported or phone_supported, f'No independently supported contact route for {host}'
            qualified = company_fit and (email_supported or phone_supported)
            if email:
                assert validate_email(email, check_deliverability=False).normalized.casefold() == email.casefold()
                if qualified:
                    all_emails.append(email.casefold())
            phone_valid = bool(phone and phonenumbers.is_valid_number(phonenumbers.parse(phone, None)))
            assert not phone or phone_valid
            assert not lead.get('email_smtp_verified')
            all_hosts.append(host)
            if phone and qualified:
                all_phones.append(phone)
            row = {'case': case, 'company': lead['company_name'], 'host': host, 'operating_company': verdict['operating_company'],
                'category_fit': verdict['category_fit'], 'location_fit': verdict['location_fit'], 'qualified': qualified,
                'contact_route_supported': bool(email_supported or phone_supported),
                'email_present': bool(email), 'email_publication_supported': email_supported,
                'mx_found': lead.get('email_mx_found', False), 'phone_present': bool(phone), 'phone_valid_format': phone_valid,
                'phone_publication_supported': phone_supported, 'direct_http_recheck': bool(observed['sources']),
                'supplemental_source_used': bool(supplemental), 'notes': verdict['notes'],
                'source_urls': [source['url'] for source in observed['sources']] + supplemental.get('urls', [])}
            case_rows.append(row)
            report_rows.append(row)
            phone_note = 'Published source reproduced' if phone_supported else 'Independent recheck inconclusive; use confirmed email route' if phone else 'No phone claimed'
            cells = [lead['company_name'], lead['company_website'], lead['location'], email, email_supported, lead.get('email_mx_found', False),
                False, phone, phone_valid, phone_note, 'Qualified company contact; delivery unconfirmed' if qualified else 'REJECTED - do not use as a prospect', verdict['notes'],
                lead.get('location_source_url', ''), lead.get('field_sources', {}).get('contact_email', ''),
                lead.get('industry_source_url', ''), lead.get('location_basis', '') + ' - see reviewed notes', lead['site_status'], lead['checked_at']]
            raw_sheet.append(cells)
            if qualified:
                sheet.append(cells)
        approved = [row for row in case_rows if row['qualified']]
        market_sheets[sheet_name] = len(approved)
        totals = {'case': case, 'requested': requested, 'returned': len(leads), 'company_fit_checked': len(case_rows),
            'qualified': len(approved), 'rejected': len(case_rows) - len(approved),
            'qualified_shortfall': requested - len(approved),
            'emails': sum(r['email_present'] for r in approved), 'email_publication_supported': sum(r['email_publication_supported'] for r in approved),
            'mx_found': sum(r['mx_found'] for r in approved), 'phones': sum(r['phone_present'] for r in approved),
            'phone_valid_format': sum(r['phone_valid_format'] for r in approved),
            'phone_publication_supported': sum(r['phone_publication_supported'] for r in approved),
            'supported_contact_routes': sum(r['contact_route_supported'] for r in approved),
            'direct_http_rechecks': sum(r['direct_http_recheck'] for r in approved),
            'supplemented_official_sources': sum(r['supplemental_source_used'] for r in approved),
            'seconds': job['result']['metadata']['elapsed_seconds']}
        stats.append(totals)
        summary_sheet.append([case, requested, len(leads), len(approved), totals['emails'], totals['mx_found'], totals['phones'],
            totals['phone_valid_format'], totals['supported_contact_routes'], totals['direct_http_rechecks'],
            totals['supplemented_official_sources'], totals['seconds']])
    assert len(all_hosts) == len(set(all_hosts)) == raw_total
    assert len(all_emails) == len(set(all_emails))
    assert len(all_phones) == len(set(all_phones))
    assert set(manual['companies']) == set(all_hosts), 'Manual manifest and delivered companies differ'
    if corrections:
        correction_sheet = workbook.create_sheet('Review corrections')
        correction_sheet.append(['Market', 'Company domain', 'Field', 'Discovery value', 'Reviewed value', 'Reason', 'Official sources'])
        for correction in corrections:
            correction_sheet.append(correction)
    notes = workbook.create_sheet('Limits and evidence')
    for row in [
        ['Item', 'Meaning'],
        ['Scope', 'Company-level prospects: business category, requested location, identity, published contact route and duplicate checks.'],
        ['Email', 'MX indicates a mail server, not a deliverable mailbox. SMTP delivery was not tested; no outreach was sent.'],
        ['Phone', 'Numbering format and source publication are separate checks. No ownership, live reachability or answer-rate claim.'],
        ['Source access', 'Some sites blocked the independent HTTP audit. Official pages were checked with a second public reader; indexed sources can be stale.'],
        ['Geography', 'Each market requires official address or business presence evidence. Any metro-area interpretation is stated in the row review notes.'],
        ['Evidence', str(evidence.resolve())],
        ['Manual decisions', str(Path(args.reviews).resolve())],
        ['Workbook layout', 'Market tabs and CSV files contain qualified prospects only. Review tabs retain every discovery result with its review decision. Original field values remain in the correction ledger and saved API responses.'],
        ['Corrections', 'Review corrections record source-based field changes separately from the preserved raw job responses.'],
        ['CSV safety', 'Leading formula characters, including + in phone numbers, are escaped with an apostrophe in CSV. XLSX stores the original values as literal text.'],
    ]:
        notes.append(row)
    for host, source in manual.get('supplemental_sources', {}).items():
        notes.append([host, ' | '.join(source.get('urls', [])) + ' | ' + source.get('note', '')])
    credit_reports = []
    for case, job in cases:
        file = root / (case + '-credits.json')
        if file.exists():
            credit_reports.append({'attempt': 'delivered run', **json.loads(file.read_text(encoding='utf-8'))})
    if args.history:
        seen_jobs = {report['job_id'] for report in credit_reports}
        for directory in args.history:
            for file in sorted(Path(directory).glob('*-credits.json')):
                credit = json.loads(file.read_text(encoding='utf-8'))
                if credit['job_id'] not in seen_jobs:
                    credit_reports.append({'attempt': 'diagnostic run: ' + Path(directory).name, **credit})
                    seen_jobs.add(credit['job_id'])
    if credit_reports:
        credits = workbook.create_sheet('Provider credits')
        credits.append(['Market', 'Attempt', 'Job ID', 'Tavily reported credits', 'Tavily API attempts', 'Unreported attempts',
            'Tavily key usage delta', 'Search credit delta', 'Extract credit delta', 'Hunter credit delta',
            'Hunter domain search attempts', 'Gemini calls', 'Reported Tavily credits per returned lead', 'Attribution'])
        requests = workbook.create_sheet('Provider request ledger')
        requests.append(['Market', 'Attempt', 'Endpoint', 'Outcome', 'HTTP status', 'Reported credits', 'Request ID'])
        cost_totals = {}
        for credit in credit_reports:
            usage = credit.get('request_usage') or {}
            tavily = usage.get('tavily') or {}
            delta = credit['account_delta'].get('tavily') or {}
            hunter = credit['account_delta'].get('hunter') or {}
            returned = credit['summary']['returned']
            credits.append([credit['case'], credit['attempt'], credit['job_id'], tavily.get('reported_credits'), tavily.get('attempts'),
                tavily.get('unreported_attempts'), delta.get('usage'), delta.get('search_usage'), delta.get('extract_usage'),
                hunter.get('credits', hunter.get('searches')), usage.get('hunter_domain_search_attempts'), usage.get('gemini_calls'),
                tavily.get('reported_credits', 0) / returned if returned else None, credit['attribution_note']])
            for record in tavily.get('records', []):
                requests.append([credit['case'], credit['attempt'], record['endpoint'], record['status'], record.get('http_status'),
                    record.get('credits'), record.get('request_id', '')])
            total = cost_totals.setdefault(credit['case'], {'runs': 0, 'tavily': 0, 'unknown': 0, 'hunter': 0, 'hunter_unknown': 0, 'gemini': 0})
            total['runs'] += 1
            total['tavily'] += tavily.get('reported_credits', 0)
            total['unknown'] += tavily.get('unreported_attempts', 0)
            total['hunter'] += hunter.get('credits', hunter.get('searches', 0))
            total['hunter_unknown'] += not bool(hunter)
            total['gemini'] += usage.get('gemini_calls', 0)
        costs = workbook.create_sheet('Total audit credits')
        costs.append(['Market', 'Runs including diagnostics', 'Tavily reported credits total', 'Tavily unreported attempts',
            'Hunter measured credits total', 'Hunter unmeasured runs', 'Gemini calls', 'Credit certainty'])
        for case, total in cost_totals.items():
            costs.append([case, total['runs'], total['tavily'], total['unknown'], total['hunter'], total['hunter_unknown'], total['gemini'],
                'Lower bound: unresolved request billing' if total['unknown'] or total['hunter_unknown'] else 'All request credits reported; Hunter measured account delta'])
        notes.append(['Credit measurement', 'Tavily request credits are provider reported, including zero at extraction billing thresholds. Unknown billing stays blank. Account deltas can include other account activity. Hunter unified credits are not added to their search/verification views.'])
        notes.append(['Tavily usage docs', 'https://docs.tavily.com/documentation/api-reference/endpoint/usage'])
        notes.append(['Hunter usage docs', 'https://hunter.io/api-documentation/v2#usage'])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='17324D')
            cell.font = Font(color='FFFFFF', bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
            if row[0].row % 2 == 0:
                for cell in row:
                    cell.fill = PatternFill('solid', fgColor='EFF5FA')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(55, max(18, len(str(column[0].value)) + 3))
    path = root / f'market-audit-{requested_total}-requested.xlsx'
    workbook.save(path)
    reopened = load_workbook(path, read_only=True)
    for name, count in market_sheets.items():
        assert reopened[name].max_row == count + 1
    for case, job in cases:
        assert reopened[('Review ' + case)[:31]].max_row == len(job['result']['leads']) + 1
    if credit_reports:
        assert reopened['Provider request ledger'].max_row == 1 + sum(len((credit.get('request_usage') or {}).get('tavily', {}).get('records', [])) for credit in credit_reports)
    assert all(cell.data_type != 'f' for sheet in reopened for row in sheet for cell in row)
    reopened.close()
    csv_sheets = list(market_sheets) + (['Provider credits', 'Total audit credits'] if credit_reports else [])
    for name in csv_sheets:
        with (root / (name + '.csv')).open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream)
            for row in workbook[name].values:
                # CSV has no cell type. Escape spreadsheet formula prefixes.
                writer.writerow(["'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value for value in row])
    result = {'cases': stats, 'total_leads': raw_total, 'requested_leads': requested_total, 'duplicate_companies': 0, 'duplicate_emails': 0, 'duplicate_phones': 0,
        'qualified_leads': sum(row['qualified'] for row in report_rows),
        'credits': credit_reports,
        'review_corrections': corrections,
        'smtp_delivery_verified': False, 'phone_reachability_verified': False, 'workbook_roundtrip': 'passed', 'records': report_rows}
    (root / 'quality-review.json').write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding='utf-8')
    print(json.dumps({'workbook': str(path), 'cases': stats, 'duplicate_companies': 0, 'workbook_roundtrip': 'passed'}), flush=True)


if __name__ == '__main__':
    main()
