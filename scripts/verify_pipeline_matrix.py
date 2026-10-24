"""Small live API-to-workbook audit. Uses configured provider credits; sends no outreach.

Start the local server, then run: python -X utf8 -m scripts.verify_pipeline_matrix
Raw results and a re-opened XLSX workbook are saved under exports/pipeline-audit.
Assertions validate the result contract; counts/shortfalls report actual discovery.
"""
import argparse
import asyncio
import json
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path

import httpx
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook

from core.services.lead_search_policy import domain_key, is_listing_source


CASES = {
    'dentists-sydney': ('Dentists', 'Sydney'),
    'lawyers-toronto': ('Law Firms', 'Toronto'),
    'architects-bengaluru': ('Architects', 'Bengaluru'),
    'solar-austin': ('Solar Installers', 'Austin, Texas'),
    'marketing-mumbai': ('Digital Marketing', 'Mumbai'),
    'accountants-london': ('Accountants', 'London, UK'),
    'saas-delhi': ('SaaS Development', 'Delhi'),
    'builders-hawaii': ('Van Upfitters / RV Builders', 'Hawaii'),
}
TERMINAL = {'completed', 'partial', 'failed', 'cancelled', 'interrupted'}


async def audit(client, name, count, output):
    industry, location = CASES[name]
    response = await client.post('/api/lead-searches', json={
        'industry': industry, 'location': location, 'target_count': count,
        'request_key': 'audit-' + uuid.uuid4().hex,
    })
    response.raise_for_status()
    job = response.json()
    last = None
    try:
        async with asyncio.timeout(720):
            while job['status'] not in TERMINAL:
                await asyncio.sleep(2)
                response = await client.get('/api/lead-searches/' + job['id'])
                response.raise_for_status()
                job = response.json()
                progress = (job['status'], len(job['result'].get('leads', [])))
                if progress != last:
                    print(json.dumps({'case': name, 'status': progress[0], 'leads': progress[1]}), flush=True)
                    last = progress
    except BaseException:
        # Do not leave a paid search orphaned if the audit is interrupted.
        await client.post('/api/lead-searches/' + job['id'] + '/cancel')
        raise
    (output / (name + '.json')).write_text(json.dumps(job, indent=2, ensure_ascii=False), encoding='utf-8')
    return check_job(job, name, count)


def check_job(job, name, count):
    industry, location = CASES[name]
    assert job['request']['target_count'] == count
    assert (job['request']['industry'], job['request']['location']) == CASES[name]
    result = job['result']
    rows = result['leads']
    meta = result['metadata']
    assert job['status'] == meta['stage']
    assert meta['completed'] is True
    assert meta['returned_count'] == len(rows)
    assert meta['shortfall'] == count - len(rows)
    assert len({domain_key(lead['company_website']) for lead in rows}) == len(rows)
    for lead in rows:
        assert lead['industry_match'] and lead['location_match'] and lead['usable'] and not lead['rejected']
        assert lead['industry_evidence'] and lead['location_evidence'] and lead['source_details']
        assert not is_listing_source(lead['company_website'], lead['industry_evidence'] + '\n' + lead['location_evidence'])
        assert not lead['contact_email'].lower().startswith('%20'), 'Encoded whitespace is not part of a mailbox'
        assert lead['contact_email'] or lead['contact_phone'] or lead['contact_form_found']
        assert not lead['email_smtp_verified'], 'MX-only discovery must not assert deliverability'
    summary = {'case': name, 'industry': industry, 'location': location, 'status': job['status'],
        'requested': count, 'returned': len(rows), 'shortfall': meta['shortfall'],
        'seconds': meta['elapsed_seconds'], 'stop_reason': meta['stop_reason'],
        'providers': meta['providers_used'], 'checked': meta['checked_targets'],
        'queries': len(meta['queries_used']), 'warnings': meta['warnings'],
        'contract_checks': 'passed', 'live_discovery': 'results_found' if rows else 'no_qualifying_results'}
    print(json.dumps(summary), flush=True)
    return summary, rows


async def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--case', action='append', choices=list(CASES))
    parser.add_argument('--count', type=int, default=2, choices=range(1, 51))
    parser.add_argument('--base-url', default='http://127.0.0.1:8000')
    parser.add_argument('--output', default='exports/pipeline-audit')
    parser.add_argument('--replay', action='store_true', help='Validate saved terminal jobs and rebuild the workbook without new provider calls')
    args = parser.parse_args()
    load_dotenv()
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=True)
    if args.replay:
        results = [check_job(json.loads((output / (name + '.json')).read_text(encoding='utf-8')), name, args.count) for name in (args.case or CASES)]
    else:
        async with httpx.AsyncClient(base_url=args.base_url, timeout=25,
                headers={'X-API-Key': os.getenv('AGENTIC_CORE_API_KEY', '')}) as client:
            results = await asyncio.gather(*(audit(client, name, args.count, output) for name in (args.case or CASES)))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Checked leads'
    columns = ['case', 'company_name', 'company_website', 'industry', 'location', 'contact_email',
               'contact_phone', 'quality_score', 'email_mx_found', 'email_smtp_verified',
               'industry_evidence', 'location_evidence', 'checked_at']
    sheet.append(columns)
    for summary, rows in results:
        for row in rows:
            # External text must remain literal even if it starts with a formula.
            sheet.append([summary['case']] + [row.get(column, '') for column in columns[1:]])
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = 's'
    path = output / 'verified-matrix.xlsx'
    workbook.save(path)
    reopened = load_workbook(path, read_only=True)
    assert reopened.active.max_row == 1 + sum(len(rows) for _, rows in results)
    reopened.close()
    report = {'checked_at': datetime.now(timezone.utc).isoformat(), 'cases': [summary for summary, _ in results],
              'workbook_roundtrip': 'passed', 'smtp_delivery_verified': False,
              'sheets_live_verified': False, 'notes': 'Sheets is tested with controlled storage in regression tests. This audit sends no outreach.'}
    (output / 'matrix-summary.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(json.dumps({'workbook': str(path), 'summary': str(output / 'matrix-summary.json'),
        'returned': sum(len(rows) for _, rows in results), 'failed_cases': [s['case'] for s, _ in results if s['status'] == 'failed']}), flush=True)
    if any(summary['status'] == 'failed' for summary, _ in results):
        raise SystemExit(1)


if __name__ == '__main__':
    asyncio.run(main())
