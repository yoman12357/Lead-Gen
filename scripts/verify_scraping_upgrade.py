"""Live API-to-workbook audit; explicitly selected cases consume provider credits."""
import argparse
import asyncio
import json
import os
import time
import uuid
from pathlib import Path

import httpx
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook
from scripts.verify_market_credits import usage_snapshot, deltas

CASES = {
    'texas-rv': {'industry': 'Van Upfitters / RV Builders', 'location': 'Texas', 'target_count': 15},
    'velit-california': {'industry': 'Van Upfitters / RV Builders', 'location': 'California', 'target_count': 5, 'mode': 'velit'},
    'accountants-london': {'industry': 'Accountants', 'location': 'London, UK', 'target_count': 5},
    'vento-grooming': {'category': 'pet_grooming', 'location': 'Austin, Texas', 'target_count': 5},
    'vento-creators': {'category': 'dog_parent_influencers', 'location': 'Seattle, WA', 'target_count': 2},
}


async def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--case', action='append', choices=CASES, required=True)
    parser.add_argument('--base-url', default='http://127.0.0.1:8001')
    parser.add_argument('--output', default='exports/scraping-upgrade')
    args = parser.parse_args()
    load_dotenv('.env')
    output = Path(args.output)
    output.mkdir(parents=True, exist_ok=True)
    async with httpx.AsyncClient(base_url=args.base_url, headers={'X-API-Key': os.getenv('AGENTIC_CORE_API_KEY', '')}, timeout=420) as client:
        for case in args.case:
            path = output / (case + '.json')
            if path.exists():
                raise RuntimeError('Refusing to repeat a paid search: ' + str(path))
            request = CASES[case]
            before = await usage_snapshot()
            started = time.monotonic()
            if 'category' in request:
                response = await client.post('/api/vento/search', json=request)
                response.raise_for_status()
                result = response.json()
                job = {'request': request, 'result': result}
            else:
                response = await client.post('/api/lead-searches', json={**request, 'request_key': str(uuid.uuid4())})
                response.raise_for_status()
                job = response.json()
                path.write_text(json.dumps(job, indent=2), encoding='utf-8')
                count = -1
                while job['status'] not in {'completed', 'partial', 'cancelled', 'failed', 'interrupted'}:
                    await asyncio.sleep(2)
                    response = await client.get('/api/lead-searches/' + job['id'])
                    response.raise_for_status()
                    job = response.json()
                    path.write_text(json.dumps(job, indent=2), encoding='utf-8')
                    if len(job['result']['leads']) != count:
                        count = len(job['result']['leads'])
                        print(json.dumps({'case': case, 'stage': job['status'], 'leads': count, 'seconds': round(time.monotonic()-started, 1)}), flush=True)
                result = job['result']
            path.write_text(json.dumps(job, indent=2, ensure_ascii=False), encoding='utf-8')
            after = await usage_snapshot()
            path.with_name(case + '-account-usage.json').write_text(json.dumps({'before': before, 'after': after, 'delta': deltas(before, after),
                'note': 'Account counters can include other activity; per-request records remain the primary attribution.'}, indent=2), encoding='utf-8')
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = 'Leads'
            columns = ['company_name','business_name','creator_name','company_website','website','contact_person_name','contact_email','email','contact_phone','phone','location','quality_score','relevance_score','industry_evidence','location_evidence','source_details','contact_sources']
            sheet.append(columns)
            for lead in result.get('leads', []):
                sheet.append([json.dumps(lead[key], ensure_ascii=False) if isinstance(lead.get(key),(dict,list)) else str(lead.get(key) or '') for key in columns])
            workbook_path = path.with_suffix('.xlsx')
            workbook.save(workbook_path)
            assert load_workbook(workbook_path).active.max_row == len(result.get('leads', [])) + 1
            meta = result.get('metadata', result)
            summary = {'case': case, 'returned': len(result.get('leads', [])), 'requested': request['target_count'], 'seconds': round(time.monotonic()-started, 1),
                       'stop_reason': meta.get('stop_reason'), 'provider_usage': meta.get('provider_usage'), 'scraping': meta.get('scraping')}
            path.with_name(case + '-summary.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
            usage = summary.get('provider_usage') or {}
            tavily = usage.get('tavily') or {}
            print(json.dumps({**summary, 'provider_usage': {'tavily_reported_credits': tavily.get('reported_credits'),
                'complete': tavily.get('complete'), 'hunter_calls': usage.get('hunter_domain_search_attempts'), 'gemini_calls': usage.get('gemini_calls')}}), flush=True)
            print(json.dumps([{'name': row.get('company_name') or row.get('business_name') or row.get('creator_name'), 'website': row.get('company_website') or row.get('website'), 'email': row.get('contact_email') or row.get('email'), 'phone': row.get('contact_phone') or row.get('phone')} for row in result.get('leads', [])]), flush=True)


if __name__ == '__main__':
    asyncio.run(main())
