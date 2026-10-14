import asyncio

from openpyxl import load_workbook

from core.services.vento.batch_runner import VentoBatchRunner, parse_locations


class FakePipeline:
    async def run(self, *, category, location, target_count):
        handle = location.lower().replace(" ", "")
        leads = [
            {
                "creator_name": f"{location} Creator",
                "instagram_handle": f"@{handle}",
                "instagram_follower_count": 76000,
                "email": f"{handle}@example.com",
                "category": category,
                "location": location,
                "niche": "Dog Parent Influencer",
                "relevance_score": 9,
                "status": "New",
                "follower_count": 75000,
            }
        ]
        if location == "Seattle":
            leads.append(
                {
                    "creator_name": "Duplicate Denver",
                    "instagram_handle": "@denver",
                    "email": "dupe@example.com",
                    "category": category,
                    "location": location,
                }
            )
        return {
            "leads": leads[:target_count],
            "metrics": {
                "raw_discovered": 3,
                "usable": len(leads[:target_count]),
                "with_email": len(leads[:target_count]),
                "with_instagram": len(leads[:target_count]),
                "tavily_calls": 2,
            },
        }


def test_parse_locations_supports_comma_newline_and_dedupes():
    assert parse_locations("Denver, Seattle\nPortland; Denver") == ["Denver", "Seattle", "Portland"]
    assert parse_locations("Seattle, WA, Portland, OR") == ["Seattle, WA", "Portland, OR"]


def test_batch_runner_writes_single_workbook_and_dedupes(tmp_path):
    output = tmp_path / "final_vento.xlsx"
    result = asyncio.run(
        VentoBatchRunner(pipeline=FakePipeline()).run(
            category="dog_parent_influencers",
            locations=["Denver", "Seattle"],
            target_per_location=2,
            output_path=str(output),
        )
    )

    assert result["returned_count"] == 2
    assert result["shortfall"] == 2
    assert result["metrics"]["same_batch_duplicates_skipped"] == 1
    assert output.exists()

    workbook = load_workbook(output)
    try:
        sheet = workbook["Vento Influencers"]
        assert sheet.max_row == 3
        headers = [cell.value for cell in sheet[1]]
        assert headers == [
            "Name (Required)",
            "Email (Required)",
            "Company Name",
            "Designation",
            "Phone",
            "Instagram",
            "TikTok",
            "Industry",
            "Address Line 1",
            "City",
            "State",
            "Country",
            "Zip",
            "Notes",
            "Tags",
            "Followers Bucket",
        ]
        assert sheet.cell(2, headers.index("Name (Required)") + 1).value == "denver"
        assert sheet.cell(2, headers.index("Designation") + 1).value == "pet influencer"
        assert sheet.cell(2, headers.index("Instagram") + 1).value == "https://www.instagram.com/denver/"
        assert sheet.cell(2, headers.index("Notes") + 1).value == "Follower Count:76000"
        assert sheet.cell(2, headers.index("Tags") + 1).value == "instagram, pet influencer"
        assert sheet.cell(2, headers.index("Followers Bucket") + 1).value == 76000
        summary = workbook["Run Summary"]
        assert summary["A1"].value == "Category"
        assert summary["B4"].value == 2
    finally:
        workbook.close()
