from __future__ import annotations

import asyncio
import re
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Set

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from core.services.vento.qualification import creator_identity_keys
from core.services.vento.vento_lead_pipeline import VentoLeadPipeline


DEFAULT_TARGET_PER_LOCATION = 20
DEFAULT_MAX_RUNTIME_SECONDS = 1800
US_STATE_CODES = {
    "al", "ak", "az", "ar", "ca", "co", "ct", "de", "fl", "ga", "hi", "id", "il", "in", "ia", "ks",
    "ky", "la", "me", "md", "ma", "mi", "mn", "ms", "mo", "mt", "ne", "nv", "nh", "nj", "nm", "ny",
    "nc", "nd", "oh", "ok", "or", "pa", "ri", "sc", "sd", "tn", "tx", "ut", "vt", "va", "wa", "wv",
    "wi", "wy", "dc",
}

VENTO_BATCH_HEADERS = [
    "Name (Required)",
    "Email (Required)",
    "Company Name",
    "Designation",
    "Phone",
    "Instagram",
    "TikTok",
    "Industry",
    "Address Line 1",
    "City",
    "State",
    "Country",
    "Zip",
    "Notes",
    "Tags",
    "Followers Bucket",
]


def parse_locations(value: Any) -> List[str]:
    """Parse comma/newline/semicolon separated location text."""
    if isinstance(value, (list, tuple)):
        raw_parts = [str(item or "") for item in value]
    else:
        raw_parts = re.split(r"[,;\n\r]+", str(value or ""))
    folded_parts: List[str] = []
    for raw in raw_parts:
        cleaned = " ".join(str(raw or "").split()).strip()
        if cleaned.lower() in US_STATE_CODES and folded_parts:
            folded_parts[-1] = f"{folded_parts[-1]}, {cleaned.upper()}"
        else:
            folded_parts.append(cleaned)

    locations: List[str] = []
    seen: Set[str] = set()
    for cleaned in folded_parts:
        key = cleaned.lower()
        if len(cleaned) < 2 or key in seen:
            continue
        seen.add(key)
        locations.append(cleaned)
    return locations


def desktop_output_path(filename: Optional[str] = None) -> Path:
    if filename:
        requested = Path(filename).expanduser()
        if requested.is_absolute() or str(requested.parent) not in {"", "."}:
            if requested.suffix.lower() != ".xlsx":
                requested = requested.with_suffix(".xlsx")
            requested.parent.mkdir(parents=True, exist_ok=True)
            return requested
    desktop = Path.home() / "Desktop"
    root = desktop if desktop.exists() and desktop.is_dir() else Path.cwd() / "exports" / "vento_batch"
    root.mkdir(parents=True, exist_ok=True)
    if not filename:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"vento_leads_final_{timestamp}.xlsx"
    clean_name = Path(filename).name
    if not clean_name.lower().endswith(".xlsx"):
        clean_name = f"{clean_name}.xlsx"
    return root / clean_name


class VentoBatchRunner:
    """Run Vento lead generation for multiple locations into one Excel file."""

    def __init__(self, *, pipeline: Optional[VentoLeadPipeline] = None) -> None:
        self.pipeline = pipeline or VentoLeadPipeline()

    async def run(
        self,
        *,
        category: str,
        locations: Sequence[str],
        target_per_location: int = DEFAULT_TARGET_PER_LOCATION,
        max_runtime_seconds: int = DEFAULT_MAX_RUNTIME_SECONDS,
        output_path: Optional[str] = None,
    ) -> Dict[str, Any]:
        clean_locations = parse_locations(list(locations))
        if not clean_locations:
            raise ValueError("At least one location is required")
        target_per_location = self._validate_target(target_per_location)
        max_runtime_seconds = self._validate_runtime(max_runtime_seconds)

        started = time.perf_counter()
        accepted: List[Dict[str, Any]] = []
        identity_keys: Set[str] = set()
        location_runs: List[Dict[str, Any]] = []
        aggregate_metrics: Dict[str, int] = {
            "raw_discovered": 0,
            "enriched": 0,
            "duplicates_skipped": 0,
            "same_batch_duplicates_skipped": 0,
            "usable": 0,
            "rejected": 0,
            "with_email": 0,
            "with_phone": 0,
            "with_instagram": 0,
            "with_tiktok": 0,
            "with_website": 0,
            "tavily_calls": 0,
            "tavily_quota_errors": 0,
            "fallback_search_calls": 0,
            "search_failures": 0,
            "search_timeouts": 0,
            "discovery_timeouts": 0,
            "enrichment_timeouts": 0,
        }

        for order, location in enumerate(clean_locations, start=1):
            remaining = max_runtime_seconds - (time.perf_counter() - started)
            if remaining <= 0:
                location_runs.append(
                    self._skipped_run(
                        location=location,
                        order=order,
                        reason="runtime_budget_exhausted",
                        target_per_location=target_per_location,
                    )
                )
                continue
            run_started = time.perf_counter()
            try:
                result = await asyncio.wait_for(
                    self.pipeline.run(
                        category=category,
                        location=location,
                        target_count=target_per_location,
                    ),
                    timeout=remaining,
                )
                leads, duplicate_count = self._append_unique_leads(
                    result.get("leads") or [],
                    accepted=accepted,
                    identity_keys=identity_keys,
                    requested_location=location,
                    batch_order=order,
                )
                metrics = dict(result.get("metrics") or {})
                metrics["same_batch_duplicates_skipped"] = duplicate_count
                self._add_metrics(aggregate_metrics, metrics)
                location_runs.append(
                    {
                        "location": location,
                        "order": order,
                        "status": "completed" if len(leads) >= target_per_location else "completed_with_shortfall",
                        "target_count": target_per_location,
                        "returned_count": len(leads),
                        "shortfall": max(0, target_per_location - len(leads)),
                        "duration_seconds": round(time.perf_counter() - run_started, 2),
                        "metrics": metrics,
                        "provider_usage": result.get('provider_usage', {}),
                    }
                )
            except asyncio.TimeoutError:
                location_runs.append(
                    self._skipped_run(
                        location=location,
                        order=order,
                        reason="runtime_budget_exhausted",
                        target_per_location=target_per_location,
                        duration_seconds=time.perf_counter() - run_started,
                    )
                )
                break

        output = desktop_output_path(output_path)
        export_stats = write_vento_batch_workbook(
            output,
            leads=accepted,
            runs=location_runs,
            category=category,
            target_per_location=target_per_location,
            max_runtime_seconds=max_runtime_seconds,
        )
        total_target = len(clean_locations) * target_per_location
        duration = round(time.perf_counter() - started, 2)
        return {
            "category": category,
            "locations": clean_locations,
            "target_per_location": target_per_location,
            "total_target": total_target,
            "returned_count": len(accepted),
            "shortfall": max(0, total_target - len(accepted)),
            "leads": accepted,
            "runs": location_runs,
            "provider_usage": {
                "tavily_reported_credits": sum((run.get('provider_usage', {}).get('tavily') or {}).get('reported_credits', 0) for run in location_runs),
                "complete": all((run.get('provider_usage', {}).get('tavily') or {}).get('complete', False) for run in location_runs),
                "per_location": [{'location': run['location'], **run.get('provider_usage', {})} for run in location_runs],
            },
            "metrics": {
                **aggregate_metrics,
                "usable": len(accepted),
                "duration_seconds": duration,
                "location_count": len(clean_locations),
                "completed_locations": sum(1 for run in location_runs if str(run.get("status", "")).startswith("completed")),
            },
            "output_file_path": str(output),
            "output_file_name": output.name,
            "export": export_stats,
            "status": "completed" if len(accepted) >= total_target else "completed_with_shortfall",
        }

    @staticmethod
    def _validate_target(value: int) -> int:
        try:
            numeric = int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError("target_per_location must be a number") from exc
        if not 1 <= numeric <= 200:
            raise ValueError("target_per_location must be between 1 and 200")
        return numeric

    @staticmethod
    def _validate_runtime(value: int) -> int:
        try:
            numeric = int(value)
        except (TypeError, ValueError) as exc:
            raise ValueError("max_runtime_seconds must be a number") from exc
        if not 60 <= numeric <= DEFAULT_MAX_RUNTIME_SECONDS:
            raise ValueError("max_runtime_seconds must be between 60 and 1800")
        return numeric

    @staticmethod
    def _identity_keys(lead: Dict[str, Any]) -> Set[str]:
        keys = set(creator_identity_keys(lead))
        try:
            keys.update(VentoLeadPipeline._dedup_keys(lead))
        except Exception:
            pass
        return {str(key).lower() for key in keys if str(key or "").strip()}

    def _append_unique_leads(
        self,
        leads: Iterable[Dict[str, Any]],
        *,
        accepted: List[Dict[str, Any]],
        identity_keys: Set[str],
        requested_location: str,
        batch_order: int,
    ) -> tuple[List[Dict[str, Any]], int]:
        appended: List[Dict[str, Any]] = []
        duplicates = 0
        for raw in leads:
            lead = dict(raw)
            keys = self._identity_keys(lead)
            if keys and keys.intersection(identity_keys):
                duplicates += 1
                continue
            lead["requested_location"] = requested_location
            lead["batch_order"] = batch_order
            lead.setdefault("status", "New")
            accepted.append(lead)
            appended.append(lead)
            identity_keys.update(keys)
        return appended, duplicates

    @staticmethod
    def _add_metrics(target: Dict[str, int], source: Dict[str, Any]) -> None:
        for key in list(target):
            try:
                target[key] += int(source.get(key) or 0)
            except (TypeError, ValueError):
                continue

    @staticmethod
    def _skipped_run(
        *,
        location: str,
        order: int,
        reason: str,
        target_per_location: int,
        duration_seconds: float = 0.0,
    ) -> Dict[str, Any]:
        return {
            "location": location,
            "order": order,
            "status": "skipped",
            "target_count": target_per_location,
            "returned_count": 0,
            "shortfall": target_per_location,
            "duration_seconds": round(duration_seconds, 2),
            "error": reason,
            "metrics": {},
        }


def write_vento_batch_workbook(
    output_path: Path,
    *,
    leads: List[Dict[str, Any]],
    runs: List[Dict[str, Any]],
    category: str,
    target_per_location: int,
    max_runtime_seconds: int,
) -> Dict[str, Any]:
    workbook = Workbook()
    lead_sheet = workbook.active
    lead_sheet.title = "Vento Influencers"
    lead_sheet.append(VENTO_BATCH_HEADERS)
    for lead in leads:
        lead_sheet.append([_cell_value(_lead_field(lead, header)) for header in VENTO_BATCH_HEADERS])
    _style_sheet(lead_sheet)

    summary = workbook.create_sheet("Run Summary")
    summary.append(["Category", category])
    summary.append(["Target Per Location", target_per_location])
    summary.append(["Max Runtime Seconds", max_runtime_seconds])
    summary.append(["Total Leads", len(leads)])
    summary.append([])
    summary.append(["Order", "Location", "Status", "Target", "Returned", "Shortfall", "Duration Seconds", "Error"])
    for run in runs:
        summary.append(
            [
                run.get("order"),
                run.get("location"),
                run.get("status"),
                run.get("target_count"),
                run.get("returned_count"),
                run.get("shortfall"),
                run.get("duration_seconds"),
                run.get("error", ""),
            ]
        )
    _style_sheet(summary)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()
    return {
        "file_path": str(output_path),
        "file_name": output_path.name,
        "lead_rows": len(leads),
        "run_rows": len(runs),
    }


def _lead_name(lead: Dict[str, Any]) -> str:
    return str(lead.get("creator_name") or lead.get("business_name") or lead.get("company_name") or lead.get("name") or "").strip()


def _lead_field(lead: Dict[str, Any], header: str) -> Any:
    if header == "Name (Required)":
        return _social_username(lead) or _lead_name(lead)
    if header == "Email (Required)":
        return lead.get("email") or ""
    if header == "Company Name":
        return lead.get("business_name") or lead.get("company_name") or _lead_name(lead)
    if header == "Designation":
        return _designation(lead)
    if header == "Phone":
        return lead.get("phone") or ""
    if header == "Instagram":
        return _social_url(lead, "instagram")
    if header == "TikTok":
        return _social_url(lead, "tiktok")
    if header == "Industry":
        return lead.get("industry") or lead.get("niche") or "Pet Industry"
    if header == "Address Line 1":
        return lead.get("address_line_1") or lead.get("address") or ""
    if header == "City":
        return lead.get("location_city") or lead.get("city") or _location_part(lead, 0)
    if header == "State":
        return lead.get("location_state") or lead.get("state") or _location_part(lead, 1)
    if header == "Country":
        return lead.get("location_country") or lead.get("country") or ""
    if header == "Zip":
        return lead.get("zip") or lead.get("zip_code") or lead.get("postal_code") or ""
    if header == "Notes":
        follower_count = _preferred_follower_count(lead)
        prefix = f"Follower Count:{follower_count}" if follower_count not in (None, "") else "Follower Count:"
        detail = lead.get("notes") or lead.get("value_proposition") or ""
        return f"{prefix} {detail}".strip()
    if header == "Tags":
        return _tags(lead)
    if header == "Followers Bucket":
        return _preferred_follower_count(lead)
    return ""


def _designation(lead: Dict[str, Any]) -> str:
    category = str(lead.get("category") or lead.get("lead_category") or "").lower()
    if category == "dog_parent_influencers" or lead.get("instagram_handle") or lead.get("tiktok_handle"):
        return "pet influencer"
    return "pet brand"


def _social_url(lead: Dict[str, Any], platform: str) -> str:
    url_key = f"{platform}_url"
    handle_key = f"{platform}_handle"
    url = str(lead.get(url_key) or "").strip()
    if url:
        return url
    handle = str(lead.get(handle_key) or "").strip().lstrip("@")
    if not handle:
        return ""
    if platform == "instagram":
        return f"https://www.instagram.com/{handle}/"
    return f"https://www.tiktok.com/@{handle}"


def _social_username(lead: Dict[str, Any]) -> str:
    for platform in ("instagram", "tiktok"):
        handle = str(lead.get(f"{platform}_handle") or "").strip().lstrip("@")
        if handle:
            return handle
        username = _username_from_url(_social_url(lead, platform), platform)
        if username:
            return username
    return ""


def _username_from_url(url: str, platform: str) -> str:
    text = str(url or "").strip()
    if not text:
        return ""
    match = re.search(r"(?:instagram\.com|tiktok\.com)/@?([^/?#\s]+)", text, re.I)
    if not match:
        return ""
    username = match.group(1).strip().strip("@")
    if platform == "instagram" and username.lower() in {"p", "reel", "stories", "explore"}:
        return ""
    return username


def _preferred_follower_count(lead: Dict[str, Any]) -> Any:
    if _social_url(lead, "instagram"):
        return lead.get("instagram_follower_count") or lead.get("follower_count") or ""
    if _social_url(lead, "tiktok"):
        return lead.get("tiktok_follower_count") or lead.get("follower_count") or ""
    return lead.get("follower_count") or ""


def _tags(lead: Dict[str, Any]) -> str:
    tags: List[str] = []
    if _social_url(lead, "instagram"):
        tags.append("instagram")
    if _social_url(lead, "tiktok"):
        tags.append("tiktok")
    tags.append(_designation(lead))
    return ", ".join(dict.fromkeys(tags))


def _location_part(lead: Dict[str, Any], index: int) -> str:
    location = str(lead.get("location") or lead.get("requested_location") or "").strip()
    parts = [part.strip() for part in location.split(",") if part.strip()]
    return parts[index] if index < len(parts) else ""


def _cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, list):
        return ", ".join(str(item) for item in value if item not in (None, ""))
    if isinstance(value, dict):
        return str(value)
    if isinstance(value, str) and value.startswith(("=", "+", "-")):
        return f"'{value}"
    return value


def _style_sheet(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for column_cells in sheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(12, min(max_length + 2, 48))
    sheet.freeze_panes = "A2"
    if sheet.max_column and sheet.max_row:
        sheet.auto_filter.ref = sheet.dimensions
